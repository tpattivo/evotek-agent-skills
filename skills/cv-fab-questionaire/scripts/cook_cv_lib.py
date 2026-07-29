"""
cook_cv_lib.py — Thư viện dùng chung cho quy trình "Cook" CV Evotek

Các hàm này lo phần formatting (merged cells, styles, borders) giống hệt template.
AI agent gọi các hàm này, truyền dữ liệu vào, không cần lo về format.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

# ═══ Styles — giống hệt Evotek_CV_sample.xlsx ═══
THIN = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))
F_TITLE = Font(name='Times New Roman', size=12, bold=True)
F_LABEL = Font(name='Times New Roman', size=12, bold=True)
F_VAL   = Font(name='Times New Roman', size=12)
A_LEFT  = Alignment(horizontal='left', vertical='center', wrap_text=True)
A_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

COL_WIDTHS = {
    'A': 7.43, 'B': 17.86, 'C': 12.43, 'D': 15.71, 'E': 13.0,
    'F': 35.43, 'G': 9.14, 'H': 14.14, 'I': 12.86, 'J': 22.14,
    'K': 36.14, 'L': 9.14, 'M': 13.0, 'N': 13.0, 'O': 13.0,
    'P': 13.0, 'Q': 13.0, 'R': 13.0, 'S': 13.0, 'T': 0.0}


# ═══ Hàm tiện ích ═══

def set_cell(ws, row, col, value, font=F_VAL, align=A_LEFT, border=THIN):
    """Gán value + style cho một cell."""
    c = ws.cell(row=row, column=col, value=value)
    c.font = font; c.alignment = align; c.border = border
    return c

def merge(ws, r1, c1, r2, c2):
    """Merge cells: (row1, col1_letter) → (row2, col2_letter)."""
    ws.merge_cells(f'{c1}{r1}:{c2}{r2}')

def row_h(ws, row, h):
    ws.row_dimensions[row].height = h

def col_widths(ws):
    for c, w in COL_WIDTHS.items():
        ws.column_dimensions[c].width = w

def apply_borders(ws, max_row, max_col):
    """Thêm thin border cho toàn bộ vùng dữ liệu."""
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = THIN


# ═══ Hàm tạo workbook trắng theo template ═══

def create_cv_workbook():
    """Tạo workbook mới với column widths, styles, và cấu trúc rỗng."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'CV mẫu'
    col_widths(ws)
    return wb, ws


# ═══ Hàm xây dựng các section ═══

def write_section_title(ws, row, title):
    """Viết section title: merged A:J, bold, border."""
    merge(ws, row, 'A', row, 'J')
    set_cell(ws, row, 1, title, F_TITLE, A_LEFT)

def write_personal_info(ws, row, label, value, row_height=27.75):
    """Viết 1 dòng thông tin cá nhân (vd: Họ và tên, Ngày sinh)."""
    merge(ws, row, 'A', row, 'B')
    set_cell(ws, row, 1, label, F_LABEL, A_LEFT)
    merge(ws, row, 'C', row, 'J')
    set_cell(ws, row, 3, value, F_VAL, A_LEFT)
    row_h(ws, row, row_height)

def write_education_entry(ws, row, stt, time, school, major, row_height=30.0):
    """Viết 1 dòng học vấn."""
    set_cell(ws, row, 1, stt, F_VAL, A_CENTER)
    set_cell(ws, row, 2, time, F_VAL, A_LEFT)
    merge(ws, row, 'C', row, 'F')
    set_cell(ws, row, 3, school, F_VAL, A_LEFT)
    merge(ws, row, 'G', row, 'J')
    set_cell(ws, row, 7, major, F_VAL, A_LEFT)
    row_h(ws, row, row_height)

def write_work_entry(ws, row, stt, start, end, company, position, row_height=28.5):
    """Viết 1 dòng quá trình công tác."""
    set_cell(ws, row, 1, stt, F_VAL, A_CENTER)
    merge(ws, row, 'B', row, 'C'); set_cell(ws, row, 2, start, F_VAL, A_LEFT)
    merge(ws, row, 'D', row, 'E'); set_cell(ws, row, 4, end, F_VAL, A_LEFT)
    merge(ws, row, 'F', row, 'H'); set_cell(ws, row, 6, company, F_VAL, A_LEFT)
    merge(ws, row, 'I', row, 'J'); set_cell(ws, row, 9, position, F_VAL, A_LEFT)
    row_h(ws, row, row_height)

def write_project(ws, label_row, stt, name, time, desc, scale, role, work, tech):
    """
    Viết 1 block dự án (7 rows), format giống template.
    Trả về row cuối của block (label_row+6).
    """
    merge(ws, label_row, 'A', label_row+6, 'A')
    set_cell(ws, label_row, 1, stt, F_VAL, A_CENTER)

    labels = ['Tên dự án', 'Thời gian tham gia dự án', 'Mô tả dự án',
              'Quy mô dự án', 'Vị trí làm việc', 'Công việc thực hiện trong dự án',
              'Công nghệ sử dụng']
    values = [name, time, desc, scale, role, work, tech]
    heights = [27.0, 27.0, 63.0, 27.0, 27.0, 174.0, 37.5]

    for i, (lbl, val) in enumerate(zip(labels, values)):
        r = label_row + i
        merge(ws, r, 'B', r, 'C'); set_cell(ws, r, 2, lbl, F_LABEL, A_LEFT)
        merge(ws, r, 'D', r, 'J'); set_cell(ws, r, 4, val, F_VAL, A_LEFT)
        row_h(ws, r, heights[i])

    return label_row + 6


# ═══ Hàm hỗ trợ update ô cho Step 2.1 (fabricate) ═══

def update_cell(ws, cell, value):
    """Ghi đè giá trị vào cell, giữ nguyên style."""
    ws[cell] = value

def update_value_cell(ws, cell, value):
    """Unmerge range chứa cell, merge lại, ghi value, áp style."""
    # Tìm merged range chứa cell này
    for mc in list(ws.merged_cells.ranges):
        if cell in mc:
            ws.unmerge_cells(str(mc))
            break
    # Merge từ cột D đến J tại row của cell
    row = int(''.join(filter(str.isdigit, cell)))
    col = ''.join(filter(str.isalpha, cell))
    # Xác định range (luôn D:J cho value cells)
    merge(ws, row, 'D', row, 'J')
    c = ws.cell(row=row, column=4)
    c.value = value
    c.font = F_VAL; c.alignment = A_LEFT; c.border = THIN
