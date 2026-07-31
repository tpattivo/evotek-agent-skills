"""
Generate a filled Evotek CV form from template and extracted candidate data.

Usage:
  Single CV:  1. Edit the DATA section below with extracted CV information.
              2. Run: python generate_form.py <path-to-original-cv-pdf>

  Parallel:   1. Write the candidate data to a JSON file (keys are the DATA
                 SECTION variable names; missing keys fall back to the
                 defaults declared below).
              2. Run: python generate_form.py <path-to-pdf> <data.json>

The output .docx will be saved in the same folder as the input PDF.
"""
import copy, json, os, sys
from docx import Document
from docx.table import Table
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║  DATA SECTION — Edit these values with the extracted CV information    ║
# ╚══════════════════════════════════════════════════════════════════════════╝

CANDIDATE_SLUG = 'CandidateSlug'  # PascalCase
OUTPUT_LANG    = 'English'

# ── Personal Info ──────────────────────────────────────────────────────────

NAME = 'Full Name'
ROLE = 'Job Title'

# ── Overview ───────────────────────────────────────────────────────────────
# Each line starts with "- " (hyphen + space). No other bullet symbols.

ACHI = (
    '- Achievement or qualification 1\n'
    '- Achievement or qualification 2\n'
    '- Achievement or qualification 3\n'
)

# ── Education ──────────────────────────────────────────────────────────────

EDU1_YEAR   = '2020 - 2024'
EDU1_DETAIL = 'Degree, University\nGPA: X.X'

EDU2_YEAR   = 'No info'
EDU2_DETAIL = 'No info'

# ── Language ───────────────────────────────────────────────────────────────

ENG_MARK = 'No info'

# ── Technical Skills ───────────────────────────────────────────────────────
# Use "No info" for any category not found in the CV.

TECH_OS      = 'No info'
TECH_DB      = 'No info'
TECH_PROG    = 'No info'
TECH_DEVTOOL = 'No info'
TECH_METHOD  = 'No info'

# ── Projects ───────────────────────────────────────────────────────────────
# Add one entry per project. Fields: name, duration, pos, size, des, res, tech
# "res" uses "- " for each bullet. Use "No info" for missing fields.

PROJECTS = [
    {
        'name': 'Project Name',
        'duration': '01/2023 - 12/2023',
        'pos': 'Role',
        'size': 'No info',
        'des': 'Project description.',
        'res': '- Responsibility 1\n- Responsibility 2',
        'tech': 'Tech 1, Tech 2',
    },
]

# ═══════════════════════════════════════════════════════════════════════════
#  END OF DATA SECTION — Do not modify below unless you understand the code
# ═══════════════════════════════════════════════════════════════════════════

INPUT_PDF = sys.argv[1] if len(sys.argv) > 1 else (
    print('Usage: python generate_form.py <path-to-pdf> [<data.json>]') or sys.exit(1)
)
INPUT_DIR = os.path.dirname(os.path.abspath(INPUT_PDF))

# ── Optional JSON data override (parallel / scripted mode) ────────────────
# Run: python generate_form.py <path-to-pdf> <data.json>
# Keys mirror the DATA SECTION variable names above; missing keys keep the
# defaults declared there.

DATA_JSON = sys.argv[2] if len(sys.argv) > 2 else None
DATA_KEYS = ['CANDIDATE_SLUG', 'OUTPUT_LANG', 'NAME', 'ROLE', 'ACHI',
             'EDU1_YEAR', 'EDU1_DETAIL', 'EDU2_YEAR', 'EDU2_DETAIL',
             'ENG_MARK', 'TECH_OS', 'TECH_DB', 'TECH_PROG', 'TECH_DEVTOOL',
             'TECH_METHOD', 'PROJECTS']

if DATA_JSON:
    with open(DATA_JSON, encoding='utf-8') as f:
        override = json.load(f)
    for key in DATA_KEYS:
        if key in override:
            globals()[key] = override[key]

# Resolve template relative to THIS script's directory
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SCRIPT_DIR, '..', 'assets', 'Evotek_cv_form_template.docx')

doc = Document(TEMPLATE_PATH)

# ── Helpers ────────────────────────────────────────────────────────────────

def set_cell(cell, text):
    """Replace cell content. Convert common bullets to hyphens."""
    text = text.replace('\u2022', '-').replace('\uf0b7', '-').replace('\u25cf', '-')
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ''
    if cell.paragraphs:
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = text
        else:
            r = OxmlElement('w:r')
            t = OxmlElement('w:t')
            t.text = text
            r.append(t)
            p._p.append(r)


def make_empty_para():
    """Create an empty w:p element (like the separators between project tables)."""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:after'), '0')
    spacing.set(qn('w:line'), '240')
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)
    p.append(pPr)
    return p


def clone_project_table(source_tbl, project_data, label):
    """
    Deep-copy source_tbl's XML, insert the clone AFTER the paragraph
    separators that follow source_tbl (so tables don't merge visually),
    then fill it with project_data. Returns the new Table wrapper.
    """
    src_elem = source_tbl._tbl
    clone_elem = copy.deepcopy(src_elem)

    # Walk forward from source table to find the last <w:p> before <w:sectPr>
    current = src_elem
    last_para = None
    while current is not None:
        tag = current.tag.split('}')[-1] if '}' in current.tag else current.tag
        if tag == 'p':
            last_para = current
        elif tag == 'sectPr':
            break
        nxt = current.getnext()
        if nxt is None:
            break
        current = nxt

    if last_para is not None:
        last_para.addnext(clone_elem)
    else:
        src_elem.addnext(clone_elem)

    # Ensure a spacing paragraph exists after the cloned table so the NEXT
    # clone (or sectPr) has a <w:p> separator to anchor to.
    nxt = clone_elem.getnext()
    nxt_tag = nxt.tag.split('}')[-1] if nxt is not None and '}' in (nxt.tag or '') else (nxt.tag if nxt is not None else '')
    if nxt is None or nxt_tag != 'p':
        clone_elem.addnext(make_empty_para())

    new_tbl = Table(clone_elem, source_tbl.part)

    # Set label in first column, first row
    set_cell(new_tbl.cell(0, 0), label)

    # Fill the 7 data rows
    fields = ['name', 'duration', 'pos', 'size', 'des', 'res', 'tech']
    for fi, fname in enumerate(fields):
        set_cell(new_tbl.cell(fi, 1), project_data[fname])

    return new_tbl


# ── Build data dict ────────────────────────────────────────────────────────

data = {
    '[[name]]': NAME,
    '[[role]]': ROLE,
    '[[achi]]': ACHI,
    '[[edu1-year]]': EDU1_YEAR,
    '[[edu1-detail]]': EDU1_DETAIL,
    '[[edu2-year]]': EDU2_YEAR,
    '[[edu2-detail]]': EDU2_DETAIL,
    '[[eng-mark]]': ENG_MARK,
    '[[tech-os]]': TECH_OS,
    '[[tech-db]]': TECH_DB,
    '[[tech-prog]]': TECH_PROG,
    '[[tech-devtool]]': TECH_DEVTOOL,
    '[[tech-method]]': TECH_METHOD,
}

# ── Fill Table 0 (left panel) ──────────────────────────────────────────────

for row in doc.tables[0].rows:
    for cell in row.cells:
        orig = cell.text
        for key, val in data.items():
            if key in orig:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if key in r.text:
                            r.text = r.text.replace(key, val)

# ── Fill built-in project tables (Tables 1, 2, 3) ──────────────────────────

# Table 1 has 8 rows: row 0 = "RECENT ASSIGNMENTS" header, rows 1-7 = Project 1
# Tables 2 & 3 have 7 rows each: row 0..6 = Project N

builtin_maps = [
    (1, 1, 0),   # (table_index, name_row_offset, projects_index)
    (2, 0, 1),
    (3, 0, 2),
]

fields = ['name', 'duration', 'pos', 'size', 'des', 'res', 'tech']

for tbl_idx, name_off, proj_idx in builtin_maps:
    if proj_idx >= len(PROJECTS):
        break
    proj = PROJECTS[proj_idx]
    for fi, fname in enumerate(fields):
        set_cell(doc.tables[tbl_idx].cell(name_off + fi, 1), proj[fname])

# ── Clone tables for remaining projects (4th, 5th, ...) ────────────────────

num_builtin = min(3, len(PROJECTS))
last_table = doc.tables[num_builtin] if num_builtin > 0 else doc.tables[0]

for pi in range(3, len(PROJECTS)):
    label = f'Project {pi + 1}'
    new_tbl = clone_project_table(last_table, PROJECTS[pi], label)
    last_table = new_tbl

# ── Save output ────────────────────────────────────────────────────────────

output_name = f'Evotek_cv_form_{CANDIDATE_SLUG}.docx'
output_path = os.path.join(INPUT_DIR, output_name)
doc.save(output_path)

print(f'Generated: {output_path}')
print(f'Tables in output: {len(doc.tables)}')
